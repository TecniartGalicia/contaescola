"""
utils/docx_balance.py
Xeración do "Balance de comedores escolares" a partir da plantilla oficial .docx.

A plantilla NON leva marcadores: as celas localízanse pola súa etiqueta (columna
esquerda) e escríbese na cela de valor conservando o formato orixinal. Así a
plantilla pode reordenarse ou retocarse sen romper o xerador.

A plantilla non vai no repositorio: vive no volume de datos (/data/plantillas/)
porque leva os nomes da dirección e o número de conta do centro.
"""
import io
import os
import re
import unicodedata

from db.connection import DB_PATH
from utils.formatters import fmt

# Ruta da plantilla: mesmo volume que a base de datos
PLANTILLA_PATH = os.environ.get(
    "PLANTILLA_COMEDOR",
    os.path.join(os.path.dirname(DB_PATH), "plantillas", "balance_comedor.docx"),
)

# Filas de gasto do impreso oficial. REPOSICIONS queda fóra por decisión do centro.
FILAS_GASTO = ["ALIMENTACION", "COMBUSTIBLE", "LIMPEZA",
               "MAQUINARIA", "OUTROS", "MANTEMENTO"]
FILA_OBVIADA = "REPOSICIONS"

# Categoría da app -> fila do impreso
CATEGORIA_A_FILA = {
    "ALIMENTACION":          "ALIMENTACION",
    "COMBUSTIBLE":           "COMBUSTIBLE",
    "LIMPEZA":               "LIMPEZA",
    "UTENSILIOS/MAQUINARIA": "MAQUINARIA",
    "UTENSILIOS":            "MAQUINARIA",   # valor herdado da migración
    "OUTROS":                "OUTROS",
    "MANTEMENTO":            "MANTEMENTO",
}


def _norm(txt: str) -> str:
    """Normaliza unha etiqueta para comparar: sen acentos, maiúsculas, sen espazos de máis."""
    t = unicodedata.normalize("NFKD", txt or "")
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t).strip().upper()


def _escribir(cela, texto: str) -> None:
    """Escribe nunha cela conservando o formato do primeiro run."""
    par = cela.paragraphs[0]
    if par.runs:
        par.runs[0].text = texto
        for r in par.runs[1:]:
            r.text = ""
    else:
        par.add_run(texto)
    for p in cela.paragraphs[1:]:      # limpar parágrafos sobrantes
        for r in p.runs:
            r.text = ""


def _fila_por_etiqueta(tabla, etiqueta: str, exacta: bool = True):
    """Devolve a fila cuxa primeira cela coincide coa etiqueta."""
    obxectivo = _norm(etiqueta)
    for fila in tabla.rows:
        actual = _norm(fila.cells[0].text)
        if (actual == obxectivo) if exacta else actual.startswith(obxectivo):
            return fila
    return None


def _set(tabla, etiqueta: str, valor: str, exacta: bool = True,
         col: int = -1, nova_etiqueta: str | None = None) -> bool:
    fila = _fila_por_etiqueta(tabla, etiqueta, exacta)
    if fila is None:
        return False
    if nova_etiqueta is not None:
        _escribir(fila.cells[0], nova_etiqueta)
    _escribir(fila.cells[col], valor)
    return True


def _eliminar_fila(tabla, etiqueta: str) -> bool:
    fila = _fila_por_etiqueta(tabla, etiqueta)
    if fila is None:
        return False
    fila._element.getparent().remove(fila._element)
    return True


def _eur(v) -> str:
    """Importe en euros; cadea baleira se non hai dato ou é cero."""
    return fmt(v) if v else ""


def gen_balance_docx(datos: dict, plantilla: str | None = None) -> bytes:
    """
    Enche a plantilla oficial cos datos calculados e devolve os bytes do .docx.

    Todo o económico vai calculado. Os datos que dependen de secretaría
    (días de funcionamento e número de alumnos por tramo) van en branco:
    complétaos a usuaria no documento descargado.
    """
    try:
        import docx
    except ImportError:
        raise RuntimeError("Falta python-docx. Executa: pip install python-docx")

    ruta = plantilla or PLANTILLA_PATH
    if not os.path.exists(ruta):
        raise FileNotFoundError(f"Non se atopa a plantilla do balance: {ruta}")

    doc = docx.Document(ruta)
    t_cab, t_alum, t_ing, t_gas, t_rem, t_fir = doc.tables[:6]

    # ── Cabeceira ────────────────────────────────────────────────
    _set(t_cab, "CURSO:", datos.get("curso", ""))
    fila_per = _fila_por_etiqueta(t_cab, "PERIODO", exacta=False)
    if fila_per is not None:
        _escribir(fila_per.cells[1], f"PERÍODO: {datos.get('periodo_txt','')}")
        _escribir(fila_per.cells[2], datos.get("trimestre", ""))
    _set(t_cab, "Días de funcionamento:", datos.get("dias", "") or "")

    # ── Datos do comedor: en branco agás o tipo ──────────────────
    _set(t_alum, "Tipo de comedor", datos.get("tipo_comedor", "C"))
    for etiqueta in ["Número de alumnos con dereito a subvención do 100%",
                     "Número de alumnos que abonan 1 € por xantar",
                     "Número de alumnos que abonan 2,5 € por xantar",
                     "Número de alumnos que abonan 4,5 € por xantar",
                     "Resto do persoal que abona 4,5 € por xantar",
                     "Persoal do comedor"]:
        _set(t_alum, etiqueta, "")

    # ── Ingresos ─────────────────────────────────────────────────
    conta = datos.get("conta", "")
    _set(t_ing, "Saldo inicial na conta da Administración", _eur(datos.get("saldo_inicial")),
         exacta=False,
         nova_etiqueta=f"Saldo inicial na conta da Administración {conta}".strip())
    _set(t_ing, "Importe das subvencións procedentes da Administración",
         _eur(datos.get("subvencions")))
    txt_oi = (datos.get("outros_ingresos_txt") or "").strip()
    _set(t_ing, "Outros ingresos", _eur(datos.get("outros_ingresos")), exacta=False,
         nova_etiqueta=f"Outros ingresos (especificar). {txt_oi}".strip()
                       if txt_oi else "Outros ingresos (especificar)")
    _set(t_ing, "Valoración das existencias no almacén a data de hoxe",
         _eur(datos.get("existencias")), exacta=False)
    _set(t_ing, "TOTAL INGRESOS", _eur(datos.get("total_ingresos")))

    # ── Gastos ───────────────────────────────────────────────────
    gastos = datos.get("gastos", {})
    _eliminar_fila(t_gas, FILA_OBVIADA)
    for fila in FILAS_GASTO:
        if fila == "OUTROS":
            txt_og = (datos.get("outros_gastos_txt") or "").strip()
            _set(t_gas, "OUTROS", _eur(gastos.get("OUTROS")), exacta=False,
                 nova_etiqueta=f"OUTROS ({txt_og})" if txt_og else "OUTROS")
        else:
            _set(t_gas, fila, _eur(gastos.get(fila)))
    _set(t_gas, "TOTAL GASTOS", _eur(datos.get("total_gastos")))

    # ── Remanente ────────────────────────────────────────────────
    _set(t_rem, "REMANENTE", _eur(datos.get("remanente")))
    _set(t_rem, "SALDO EXISTENTE A", _eur(datos.get("remanente")), exacta=False,
         nova_etiqueta=f"SALDO EXISTENTE A {datos.get('data_saldo','')}")

    # ── Lugar e data da sinatura ─────────────────────────────────
    lugar_data = datos.get("lugar_data", "")
    if lugar_data:
        for fila in t_fir.rows:
            for cela in fila.cells:
                if _norm(cela.text).startswith(_norm(datos.get("lugar", "")) + ",") or \
                   re.match(r"^[A-Za-zÁÉÍÓÚÑáéíóúñ ]+, a \d", cela.text.strip()):
                    _escribir(cela, lugar_data)
                    break

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
